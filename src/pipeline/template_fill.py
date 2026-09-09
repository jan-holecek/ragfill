from pathlib import Path
from typing import Generator

import openpyxl

from config import SearchSettings
from db.elasticsearch import ElasticSearchDB
from models.response import RAGResponse
from pipeline.embedding import Embedding
from pipeline.generation import Generation
from pipeline.query_rewrite import QueryRewrite
from docx import Document
from docx.oxml.ns import qn
import zipfile
import xml.etree.ElementTree as ET
import re
import time
import json
from pipeline.search import Search
from prompts import build_template_fill_prompt, build_template_fill_all_prompt
from models.response import TemplateFillResponse, PlaceholderResponse
import io

PLACEHOLDER_RE = re.compile(r"\{\{[^}]+\}\}")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

class TemplateFill:
    def __init__(self, db: ElasticSearchDB, embedding: Embedding, generation: Generation, rewriter: QueryRewrite | None = None) -> None:
        self.embedding = embedding
        self.search = Search(db, SearchSettings(
            bm25_K=1,
            knn_K=1,
            num_candidates=100,
            decomposed_bm25_K=1,
            decomposed_knn_K=1,
            max_total_chunks=3,
        ))
        self.generation = generation
        self.rewriter = rewriter

    def _extract(self, content: bytes, filename: str) -> dict[str, str]:
        suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if suffix == "docx":
            return self._extract_docx(content)
        elif suffix == "xlsx":
            return self._extract_xlsx(content)

        raise ValueError(f"unsupported file type '.{suffix}' (expected .docx or .xlsx)")

    def _extract_xlsx(self, content: bytes) -> dict[str, str]:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        placeholders = {}

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or cell.comment is None:
                        continue

                    for match in PLACEHOLDER_RE.findall(cell.value):
                        placeholders[match] = cell.comment.text

        return placeholders

    def _extract_docx(self, content: bytes) -> dict[str, str]:
        placeholders = {}

        with zipfile.ZipFile(io.BytesIO(content), "r") as zip_file:
            if "word/comments.xml" not in zip_file.namelist():
                return {}

            comments_xml = zip_file.read("word/comments.xml")
            root = ET.fromstring(comments_xml)
            comments = {}

            for comment in root.findall("w:comment", NS):
                comment_id = comment.get(f"{{{NS['w']}}}id")
                comment_text = "".join(text.text or "" for text in comment.findall(".//w:t", NS))
                comments[comment_id] = comment_text

        doc = Document(io.BytesIO(content))
        body = doc.element.body

        for elem in body.iter():
            if elem.tag == qn("w:commentRangeStart"):
                comment_id = elem.get(qn("w:id"))

                if comment_id in comments:
                    parent = elem.getparent()

                    if parent is not None:
                        text = "".join(t.text or "" for t in parent.findall(f".//{qn('w:t')}"))
                        matches = re.findall(r"\{\{[^}]+\}\}", text)

                        for match in matches:
                            placeholders[match] = comments[comment_id]

        return placeholders

    def _save_xlsx(self, file_path: str, output_path: str, values: dict[str, str]) -> None:
        workbook = openpyxl.load_workbook(file_path)

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue

                    for match in PLACEHOLDER_RE.findall(cell.value):
                        if match in values:
                            cell.value = cell.value.replace(match, values[match])

        workbook.save(output_path)

    def _save_docx(self, file_path: str, output_path: str, values: dict[str, str]) -> None:
        doc = Document(file_path)

        for paragraph in doc.paragraphs:
            for match in PLACEHOLDER_RE.findall(paragraph.text):
                if match in values:
                    paragraph.text = paragraph.text.replace(match, values[match])

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for match in PLACEHOLDER_RE.findall(cell.text):
                        if match in values:
                            cell.text = cell.text.replace(match, values[match])

        doc.save(output_path)

    def _generate_placeholder_content(self, prompt: str) -> RAGResponse | None:
        rewrite_result = None
        rewritten_prompt = prompt

        if self.rewriter:
            rewrite_result = self.rewriter.rewrite(prompt)
            rewritten_prompt = rewrite_result.rewritten_query

        query_vector = self.embedding.embed_query(rewritten_prompt)
        search_chunks, used_queries = self.search.rewrite_and_search(rewritten_prompt, self.embedding)

        if not search_chunks:
            return None

        response = self.generation.generate(
            prompt,
            search_chunks,
            query_vector,
            rewrite_result,
            prompt_builder=build_template_fill_prompt
        )
        response.used_queries = used_queries

        return response

    def _generate_all_placeholders(self, placeholders: dict[str, str]):
        all_chunks = {}
        used_queries_by_placeholder = {}

        for placeholder, prompt in placeholders.items():
            results, used_queries = self.search.rewrite_and_search(prompt, self.embedding)
            used_queries_by_placeholder[placeholder] = used_queries

            for chunk in results:
                all_chunks[chunk.id] = chunk

        if not all_chunks:
            return {p: "" for p in placeholders}, 0, None, used_queries_by_placeholder

        chunks = list(all_chunks.values())
        fields_description = "\n".join([f'"{p}": "{prompt}"' for p, prompt in placeholders.items()])
        llm_instruction = f"Fill the document template placeholders based on the provided context according to these requested fields:\n{fields_description}"

        response = self.generation.generate(
            query=llm_instruction,
            chunks=chunks,
            prompt_builder=build_template_fill_all_prompt
        )

        content = response.answer.strip()
        content = re.sub(r"```json\s*|\s*```", "", content).strip()
        try:
            values = json.loads(content)

            return {k: str(v) for k, v in values.items()}, response.elapsed, response, used_queries_by_placeholder
        except Exception:
            return {p: "" for p in placeholders}, response.elapsed, None, used_queries_by_placeholder

    def get_placeholders_results_stream(self, file_path: str) -> Generator[PlaceholderResponse, None, None]:
        with open(file_path, "rb") as f:
            placeholders = self._extract(f.read(), file_path)

        for placeholder, prompt in placeholders.items():
            generated = self._generate_placeholder_content(prompt)

            if generated:
                yield PlaceholderResponse(
                    placeholder=placeholder,
                    prompt=prompt,
                    answer=generated.answer,
                    completion_tokens=generated.completion_tokens,
                    prompt_tokens=generated.prompt_tokens,
                    elapsed=generated.elapsed,
                    rewrite=generated.rewrite,
                    chunks=generated.chunks,
                    used_queries=generated.used_queries,
                )
            else:
                yield PlaceholderResponse(
                    placeholder=placeholder,
                    prompt=prompt,
                    answer="",
                    completion_tokens=0,
                    prompt_tokens=0,
                    elapsed=0,
                )

    def get_placeholders_results(self, file_path: str) -> TemplateFillResponse:
        with open(file_path, "rb") as f:
            placeholders = self._extract(f.read(), file_path)
        start = time.time()
        values, elapsed, response, used_queries_by_placeholder = self._generate_all_placeholders(placeholders)

        results = [
            PlaceholderResponse(
                placeholder=p,
                prompt=placeholders[p],
                answer=values.get(p, ""),
                completion_tokens=response.completion_tokens if response else 0,
                prompt_tokens=response.prompt_tokens if response else 0,
                elapsed=elapsed,
                chunks=response.chunks if response else [],
                rewrite=response.rewrite if response else None,
                used_queries=used_queries_by_placeholder.get(p),
            )

            for p in placeholders
        ]

        return TemplateFillResponse(
            results=results,
            total_elapsed=round(time.time() - start, 2),
            total_completion_tokens=response.completion_tokens if response else 0,
            total_prompt_tokens=response.prompt_tokens if response else 0,
        )

    def save(self, file_path: str, output_path: str, values: dict[str, str]) -> None:
        suffix = file_path.rsplit(".", 1)[-1].lower() if "." in file_path else ""

        if suffix == "docx":
            self._save_docx(file_path, output_path, values)
        elif suffix == "xlsx":
            self._save_xlsx(file_path, output_path, values)
        else:
            raise ValueError(f"unsupported file type '.{suffix}' (expected .docx or .xlsx)")