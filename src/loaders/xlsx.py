import openpyxl
from pathlib import Path
from typing import Iterator
from langchain_core.documents import Document
from core.loader import BaseLoader

class XLSXLoader(BaseLoader):
    def __init__(self, file_path: str):
        super().__init__(file_path=file_path, extensions=[".xlsx"])

    def lazy_load(self) -> Iterator[Document]:
        workbook = openpyxl.load_workbook(str(self.file_path), data_only=True)
        properties = workbook.properties
        content = []

        for sheet in workbook.worksheets:
            rows = []

            for row in sheet.iter_rows(values_only=True):
                rows.append([str(cell) if cell is not None else "" for cell in row])

            rows = [r for r in rows if any(c.strip() for c in r)]

            if not rows:
                continue

            content.append(f"## {sheet.title}")

            header = rows[0]
            content.append("| " + " | ".join(header) + " |")
            content.append("| " + " | ".join(["---"] * len(header)) + " |")

            for row in rows[1:]:
                content.append("| " + " | ".join(row[:len(header)]) + " |")

        metadata = {
            "source": Path(self.file_path).name,
            "file_size": Path(self.file_path).stat().st_size,
            "word_count": len(" ".join(content).split()),
            "sheet_count": len(workbook.worksheets),
        }

        for attribute in ["title", "subject", "creator", "description", "keywords", "category", "lastModifiedBy", "version", "identifier"]:
            value = getattr(properties, attribute, None)

            if value:
                metadata[attribute.lower()] = str(value)

        for attribute in ["created", "modified", "lastPrinted"]:
            value = getattr(properties, attribute, None)

            if value:
                metadata[attribute.lower()] = value.isoformat()

        yield Document(page_content="\n".join(content), metadata=metadata)